"""A planilha de trabalho, em XLSX.

Só existe uma coluna que é sua: `acao`. Escreva `aceitar` ou `rejeitar` nela, salve, e o `aplicar`
leva para o Parsifal. A automação nunca escreve nessa coluna - é o que garante que rodar
`classificar` de novo não desfaz decisão sua.

O formato é XLSX e não CSV porque o CSV aberto no editor perdia o que era salvo, e porque a planilha
traz filtro e busca por coluna, que é o que falta na tela do Parsifal.
"""

from __future__ import annotations

import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ACEITAR = "aceitar"
REJEITAR = "rejeitar"
SEM_ACAO = ""

# Escritas que valem como decisão. Qualquer outra coisa na coluna vira erro, não silêncio.
SINONIMOS = {
    "aceitar": ACEITAR,
    "aceito": ACEITAR,
    "aceita": ACEITAR,
    "accept": ACEITAR,
    "rejeitar": REJEITAR,
    "rejeitado": REJEITAR,
    "rejeitada": REJEITAR,
    "reject": REJEITAR,
}

COLUNAS = (
    ("id", 12),
    ("acao", 12),
    ("criterio", 60),
    ("titulo", 70),
    ("sugestao", 12),
    ("sinais", 60),
    ("notas_etapa_2", 40),
    ("recorte_em", 24),
    ("doi", 28),
)

COLUNA_DA_ACAO = 2


def normalizar_acao(valor: str | None) -> str:
    """Devolve `aceitar`, `rejeitar` ou vazio. Levanta ValueError no que não for nenhum dos três."""
    texto = str(valor or "").strip().lower()
    if not texto:
        return SEM_ACAO

    sem_acento = "".join(
        letra for letra in unicodedata.normalize("NFD", texto) if unicodedata.category(letra) != "Mn"
    )
    if sem_acento not in SINONIMOS:
        raise ValueError(f"ação não reconhecida: {valor!r}. Use 'aceitar' ou 'rejeitar'.")

    return SINONIMOS[sem_acento]


def escrever(caminho: Path, linhas: list[dict[str, str]], criterios: list[str] | None = None) -> None:
    """Grava a planilha com filtro, painel congelado e lista suspensa na coluna de ação."""
    livro = Workbook()
    aba = livro.active
    aba.title = "triagem"

    for numero, (nome, largura) in enumerate(COLUNAS, start=1):
        celula = aba.cell(row=1, column=numero, value=nome)
        celula.font = Font(bold=True)
        aba.column_dimensions[get_column_letter(numero)].width = largura

    for linha_numero, linha in enumerate(linhas, start=2):
        for coluna_numero, (nome, _) in enumerate(COLUNAS, start=1):
            aba.cell(row=linha_numero, column=coluna_numero, value=linha.get(nome, ""))

    ultima = max(len(linhas) + 1, 2)
    aba.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ultima}"
    aba.freeze_panes = "A2"

    letra_da_acao = get_column_letter(COLUNA_DA_ACAO)
    escolha_de_acao = DataValidation(
        type="list", formula1=f'"{ACEITAR},{REJEITAR}"', allow_blank=True, showDropDown=False
    )
    aba.add_data_validation(escolha_de_acao)
    escolha_de_acao.add(f"{letra_da_acao}2:{letra_da_acao}{ultima}")

    if criterios:
        aba_criterios = livro.create_sheet("criterios")
        aba_criterios["A1"] = "critérios cadastrados no Parsifal"
        aba_criterios["A1"].font = Font(bold=True)
        aba_criterios.column_dimensions["A"].width = 110
        for numero, texto in enumerate(criterios, start=2):
            aba_criterios.cell(row=numero, column=1, value=texto).alignment = Alignment(wrap_text=True)

    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro.save(caminho)


def ler(caminho: Path) -> list[dict[str, str]]:
    livro = load_workbook(caminho, data_only=True)
    aba = livro["triagem"] if "triagem" in livro.sheetnames else livro.active

    cabecalho = [str(celula.value or "") for celula in aba[1]]
    linhas = []
    for valores in aba.iter_rows(min_row=2, values_only=True):
        linha = {nome: ("" if valor is None else str(valor)) for nome, valor in zip(cabecalho, valores)}
        if linha.get("id", "").strip():
            linhas.append(linha)

    return linhas
